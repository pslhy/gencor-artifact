from datetime import timedelta
import regex as re
import glob
import os
from openpyxl import load_workbook
def task_func(directory_path='./xlsx_files/'):
    count = 0
    for file_path in glob.glob(os.path.join(directory_path, '*.xlsx')):
        wb = load_workbook(file_path)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if isinstance(cell, str):
                        new_value = re.sub(r'"', r'\"', cell)
                        cell_addr = f'{cell[0]}{cell[1]}'
                        sheet[cell_addr] = new_value
        wb.save(file_path)
        count += 1
    return count


# <System Prompt>
#     You are an expert in Python programming. Solve the given problem by writing solution code in Python. When answering, insert the solution code in a ```python...``` block.
#     Do not include test cases or explanations as comments in the code.
# 
# < User Prompt >
#     Protects all double quotes in all Excel (.xlsx) files in the specified directory by prefixing them with a double backslash.
#     The function should output with:
#         int: The number of Excel files processed.
#     You should write self-contained code starting with:
#     ```
#     import regex as re
#     import glob
#     import os
#     from openpyxl import load_workbook
#     def task_func(directory_path='./xlsx_files/'):
#     ```
