# <System Prompt>
#     You are an expert in Python programming. Solve the given problem by writing solution code in Python. When answering, insert the solution code in a ```python...``` block.
#     Do not include test cases or explanations as comments in the code.
# 
# 
# === Turn 1 ===
# < User Prompt >
#     Protects all double quotes in all Excel (.xlsx) files in the specified directory by prefixing them with a double backslash.
#     The function should output with:
#         int: The number of Excel files processed.
#     You should write self-contained code starting with:
#     ```
#     import regex as re
#     import glob
#     import os
#     from openpyxl import load_workbook
#     def task_func(directory_path='./xlsx_files/'):
#     ```
from datetime import timedelta
import regex as re
import glob
import os
from openpyxl import load_workbook
def task_func(directory_path='./xlsx_files/'):
    processed_files = 0
    file_pattern = os.path.join(directory_path, '*.xlsx')
    for file_path in glob.glob(file_pattern):
        try:
            workbook = load_workbook(file_path)
            for sheet in workbook.sheetnames:
                worksheet = workbook[sheet]
                for row in worksheet.iter_rows(values_only=True):
                    for cell in row:
                        if isinstance(cell, str):
                            cell_value = re.sub(r'(?<!\\)\"', r'\\"', cell)
                            worksheet.cell(row=worksheet._cell_from劲放着异步的回调函数 str(row) + 1, column=col + 1, value=cell_value)
                workbook.save(file_path)
            processed_files += 1
        except Exception as e:
            print(f"Error processing {file_path}: {e}")
    return processed_files

# === Turn 2 ===
# < User Prompt >
#     The following errors were found in the previous code:
#     - <python>:18:48: invalid syntax. Perhaps you forgot a comma?
#     Please regenerate the code, taking these errors into account.
from datetime import timedelta
import regex as re
import glob
import os
from openpyxl import load_workbook
def task_func(directory_path='./xlsx_files/'):
    processed_files = 0
    file_pattern = os.path.join(directory_path, '*.xlsx')
    for file_path in glob.glob(file_pattern):
        try:
            workbook = load_workbook(file_path)
            for sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]
                for row in worksheet.iter_rows(values_only=True):
                    for cell in row:
                        if isinstance(cell, str):
                            cell_value = re.sub(r'(?<!\\)\"', r'\\"', cell)
                            worksheet.cell(row=row[0] + 1, column=cell[0], value=cell_value)
                workbook.save(file_path)
            processed_files += 1
        except Exception as e:
            print(f"Error processing {file_path}: {e}")
    return processed_files

