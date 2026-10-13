from datetime import timedelta
import regex as re
import glob
import os
from openpyxl import load_workbook
def task_func(directory_path='./xlsx_files/'):
    glob_path = os.path.join(directory_path, '*.xlsx')
    xlsx_files = glob.glob(glob_path)
    count = 0
    for file_name in xlsx_files:
        print(f"Processing file: {file_name}")
        wb = load_workbook(file_name)
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    if cell.data_type == 's':
                        cell.value = re.sub(r'"', r'\"', cell.value)
        wb.save(file_name)
        count += 1
    return count


# <System Prompt>
#     You are an expert in Python programming. Solve the given problem by writing solution code in Python. When answering, insert the solution code in a ```python...``` block.
#     Do not include test cases or explanations as comments in the code.
# 
# < User Prompt >
#     Protects all double quotes in all Excel (.xlsx) files in the specified directory by prefixing them with a double backslash.
#     The function should output with:
#         int: The number of Excel files processed.
#     You should write self-contained code starting with:
#     ```
#     import regex as re
#     import glob
#     import os
#     from openpyxl import load_workbook
#     def task_func(directory_path='./xlsx_files/'):
#     ```
