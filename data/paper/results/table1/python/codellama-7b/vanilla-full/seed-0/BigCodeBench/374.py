from datetime import timedelta
import regex as re
import glob
import os
from openpyxl import load_workbook
def task_func(directory_path='./xlsx_files/'):

    excel_files = glob.glob(os.path.join(directory_path, '*.xlsx'))
    num_files_processed = 0

    for excel_file in excel_files:
        wb = load_workbook(excel_file)
        for sheet in wb.worksheets:
            rows = sheet.iter_rows()
            for row in rows:
                for cell in row:
                    if cell.value is not None:
                        cell.value = re.sub(r'"', '\\"', cell.value)

        num_files_processed += 1
        wb.save(excel_file)

    return num_files_processed


# <System Prompt>
#     You are an expert in Python programming. Solve the given problem by writing solution code in Python. When answering, insert the solution code in a ```python...``` block.
#     Do not include test cases or explanations as comments in the code.
# 
# < User Prompt >
#     Protects all double quotes in all Excel (.xlsx) files in the specified directory by prefixing them with a double backslash.
#     The function should output with:
#         int: The number of Excel files processed.
#     You should write self-contained code starting with:
#     ```
#     import regex as re
#     import glob
#     import os
#     from openpyxl import load_workbook
#     def task_func(directory_path='./xlsx_files/'):
#     ```
