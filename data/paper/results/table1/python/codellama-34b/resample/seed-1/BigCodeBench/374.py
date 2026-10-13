# <System Prompt>
#     You are an expert in Python programming. Solve the given problem by writing solution code in Python. When answering, insert the solution code in a ```python...``` block.
#     Do not include test cases or explanations as comments in the code.
# 
# 
# === Turn 1 ===
# < User Prompt >
#     Protects all double quotes in all Excel (.xlsx) files in the specified directory by prefixing them with a double backslash.
#     The function should output with:
#         int: The number of Excel files processed.
#     You should write self-contained code starting with:
#     ```
#     import regex as re
#     import glob
#     import os
#     from openpyxl import load_workbook
#     def task_func(directory_path='./xlsx_files/'):
#     ```
from datetime import timedelta
import regex as re
import glob
import os
from openpyxl import load_workbook
def task_func(directory_path='./xlsx_files/'):
    file_count = 0
    for filename in glob.glob(os.path.join(directory_path, '*.xlsx')):
        book = load_workbook(filename)
        
        sheet_names = book.sheetnames
        for sheet_name in sheet_names:
            sheet = book.get_sheet_by_name(sheet_name)
            for row_cells in sheet.iter_rows(values_only=True):
                for cell in row_cells:
                    if cell.data_type =='s':
                        cell.value = re.sub(r'"', r'\"', cell.value)
            
        book.save(filename)
        file_count += 1
    
    return file_count

